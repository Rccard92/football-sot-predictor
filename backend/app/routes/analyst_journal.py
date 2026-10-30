"""Diario analisi (dal 30/10/2026): giocate, scheda pronostico, progressioni di cassa e report.

Chi vede cosa (decisione dell'utente, 30/10/2026): un analista vede e modifica solo le sue giocate e
progressioni; gli amministratori vedono tutto (in sola lettura per i diari degli altri).
Anche gli amministratori possono scrivere nel proprio diario (richiesta dell'utente, 30/10/2026): chi entra
con l'accesso di emergenza scrive a nome di un account amministratore con l'email di emergenza, creato al
primo salvataggio se non esiste ancora.
"""

from __future__ import annotations

import dataclasses
import io
from datetime import date, datetime, timezone
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.access_control import require_principal
from app.core.config import get_settings
from app.core.database import get_db
from app.core.user_auth import KIND_EMERGENCY_ADMIN, Principal, generate_password, hash_password, normalize_email
from app.models.analyst_journal import CYCLE_CLOSED, CYCLE_OPEN, ENTRY_PENDING, AnalystCycle, AnalystEntry
from app.models.app_user import ROLE_ADMIN, AppUser
from app.models.cecchino_today_fixture import CecchinoTodayFixture
from app.services.analyst.journal import (
    CHECKLIST,
    MARKET_BY_KEY,
    MARKETS,
    SOURCES,
    build_report,
    checklist_hints,
    clean_checklist,
    cycle_entries,
    cycle_payload,
    entry_payload,
    fixture_fields,
    is_locked,
    kpi_row,
    settle_entries,
    tool_snapshot,
)
from app.services.analyst.staking import DEFAULTS, MAX_BETS, METHOD_LABELS, METHODS, normalize_settings

router = APIRouter(prefix="/analyst-journal", tags=["analyst-journal"])


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _scope(principal: Principal, user_id: int | None) -> int | None:
    """Di chi sono i dati richiesti: l'analista vede solo i suoi; l'amministratore tutti o uno a scelta."""
    if principal.is_admin:
        return user_id
    if principal.user is None:
        raise HTTPException(status_code=403, detail="Serve un account")
    return principal.user.id


def _emergency_account(db: Session) -> AppUser | None:
    email = normalize_email(get_settings().emergency_admin_email)
    if "@" not in email:
        return None
    user = db.execute(select(AppUser).where(AppUser.email == email)).scalar_one_or_none()
    return user if user is None or user.is_active else None


def journal_principal(principal: Principal = Depends(require_principal), db: Session = Depends(get_db)) -> Principal:
    """L'accesso di emergenza usa l'account con l'email di emergenza, se esiste (per vedere e modificare il suo diario)."""
    if principal.kind == KIND_EMERGENCY_ADMIN and principal.user is None:
        user = _emergency_account(db)
        if user is not None:
            return dataclasses.replace(principal, user=user)
    return principal


def _writer(principal: Principal, db: Session) -> AppUser:
    if principal.user is not None:
        return principal.user
    if principal.kind == KIND_EMERGENCY_ADMIN:
        email = normalize_email(get_settings().emergency_admin_email)
        if "@" in email and db.execute(select(AppUser.id).where(AppUser.email == email)).first() is None:
            # Password a caso mai mostrata: si entra sempre con la password di emergenza
            # (o la si reimposta dal pannello utenti per farne un account normale).
            user = AppUser(
                email=email,
                first_name="Amministratore",
                last_name="(emergenza)",
                role=ROLE_ADMIN,
                password_hash=hash_password(generate_password()),
                must_change_password=True,
                is_active=True,
            )
            db.add(user)
            db.flush()
            return user
    raise HTTPException(status_code=403, detail="Per scrivere nel diario serve un account attivo")


def _own_entry(db: Session, principal: Principal, entry_id: int) -> AnalystEntry:
    entry = db.get(AnalystEntry, entry_id)
    if entry is None:
        raise HTTPException(status_code=404, detail="Giocata non trovata")
    if principal.user is None or entry.user_id != principal.user.id:
        raise HTTPException(status_code=403, detail="Puoi modificare solo le tue giocate")
    return entry


def _own_cycle(db: Session, principal: Principal, cycle_id: int) -> AnalystCycle:
    cycle = db.get(AnalystCycle, cycle_id)
    if cycle is None:
        raise HTTPException(status_code=404, detail="Progressione non trovata")
    if principal.user is None or cycle.user_id != principal.user.id:
        raise HTTPException(status_code=403, detail="Puoi modificare solo le tue progressioni")
    return cycle


def _visible_cycle(db: Session, principal: Principal, cycle_id: int) -> AnalystCycle:
    cycle = db.get(AnalystCycle, cycle_id)
    if cycle is None:
        raise HTTPException(status_code=404, detail="Progressione non trovata")
    if not principal.is_admin and (principal.user is None or cycle.user_id != principal.user.id):
        raise HTTPException(status_code=403, detail="Progressione di un altro utente")
    return cycle


def _names(db: Session, ids: set[int]) -> dict[int, str]:
    if not ids:
        return {}
    rows = db.execute(select(AppUser.id, AppUser.first_name, AppUser.last_name).where(AppUser.id.in_(ids))).all()
    return {uid: f"{fn} {ln}" for uid, fn, ln in rows}


def _float(value: Any, field: str, *, minimum: float | None = None) -> float | None:
    if value is None or value == "":
        return None
    try:
        out = float(value)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail=f"{field}: numero non valido") from None
    if minimum is not None and out < minimum:
        raise HTTPException(status_code=400, detail=f"{field}: deve essere almeno {minimum:g}")
    return out


def _entries_query(principal: Principal, *, user_id, date_from, date_to, cycle_id, outcome):
    q = select(AnalystEntry)
    scope = _scope(principal, user_id)
    if scope is not None:
        q = q.where(AnalystEntry.user_id == scope)
    if date_from:
        q = q.where(AnalystEntry.match_date >= date_from)
    if date_to:
        q = q.where(AnalystEntry.match_date <= date_to)
    if cycle_id:
        q = q.where(AnalystEntry.cycle_id == cycle_id)
    if outcome:
        q = q.where(AnalystEntry.outcome == outcome)
    return q.order_by(AnalystEntry.match_date.desc(), AnalystEntry.kickoff.desc().nulls_last(), AnalystEntry.id.desc())


# --- informazioni fisse ---------------------------------------------------------------------------


@router.get("/meta")
def meta(principal: Principal = Depends(journal_principal)) -> dict[str, Any]:
    return {
        "markets": MARKETS,
        "checklist": CHECKLIST,
        "sources": SOURCES,
        "methods": [
            {"key": m, "label": METHOD_LABELS[m], "defaults": DEFAULTS[m], "max_bets": MAX_BETS[m]} for m in METHODS
        ],
        "can_write": principal.user is not None or principal.kind == KIND_EMERGENCY_ADMIN,
        "is_admin": principal.is_admin,
        "user_id": principal.user.id if principal.user else None,
    }


@router.get("/analysts")
def analysts(principal: Principal = Depends(journal_principal), db: Session = Depends(get_db)) -> dict[str, Any]:
    """Per l'amministratore: chi ha un diario (per il filtro)."""
    if not principal.is_admin:
        me = principal.user
        return {"analysts": [{"id": me.id, "name": f"{me.first_name} {me.last_name}"}] if me else []}
    ids = {uid for (uid,) in db.execute(select(AnalystEntry.user_id).distinct()).all()}
    ids |= {uid for (uid,) in db.execute(select(AnalystCycle.user_id).distinct()).all()}
    names = _names(db, ids)
    return {"analysts": sorted(({"id": i, "name": n} for i, n in names.items()), key=lambda x: x["name"])}


# --- partita (per il dialogo "Aggiungi al diario") ------------------------------------------------


@router.get("/fixture/{today_fixture_id}")
def fixture_info(
    today_fixture_id: int,
    market_key: str | None = None,
    principal: Principal = Depends(journal_principal),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    fx = db.get(CecchinoTodayFixture, today_fixture_id)
    if fx is None:
        raise HTTPException(status_code=404, detail="Partita non trovata")
    fields = fixture_fields(fx)
    kickoff = fx.kickoff if fx.kickoff is None or fx.kickoff.tzinfo else fx.kickoff.replace(tzinfo=timezone.utc)
    markets = []
    for m in MARKETS:
        row = kpi_row(fx, m["key"]) or {}
        markets.append({**m, "quota_book": row.get("quota_book"), "prob_cecchino": row.get("prob_cecchino")})
    mine = []
    if principal.user is not None:
        mine = [
            entry_payload(e)
            for e in db.execute(
                select(AnalystEntry).where(
                    AnalystEntry.user_id == principal.user.id, AnalystEntry.today_fixture_id == today_fixture_id
                )
            ).scalars()
        ]
    return {
        "fixture": {
            **fields,
            "match_date": fields["match_date"].isoformat(),
            "kickoff": kickoff.isoformat() if kickoff else None,
            "started": bool(kickoff and kickoff <= _now()),
        },
        "markets": markets,
        "hints": checklist_hints(db, fx, market_key),
        "my_entries": mine,
    }


# --- giocate --------------------------------------------------------------------------------------


@router.get("/entries")
def list_entries(
    date_from: date | None = None,
    date_to: date | None = None,
    user_id: int | None = None,
    cycle_id: int | None = None,
    outcome: str | None = None,
    principal: Principal = Depends(journal_principal),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    settle_entries(db, user_id=_scope(principal, user_id))
    rows = db.execute(
        _entries_query(principal, user_id=user_id, date_from=date_from, date_to=date_to, cycle_id=cycle_id, outcome=outcome).limit(2000)
    ).scalars().all()
    names = _names(db, {e.user_id for e in rows})
    now = _now()
    return {"entries": [entry_payload(e, user_name=names.get(e.user_id), now=now) for e in rows]}


@router.post("/entries")
def create_entry(body: dict[str, Any], principal: Principal = Depends(journal_principal), db: Session = Depends(get_db)) -> dict[str, Any]:
    user = _writer(principal, db)
    body = body or {}
    try:
        fixture_id = int(body.get("today_fixture_id"))
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Partita mancante") from None
    fx = db.get(CecchinoTodayFixture, fixture_id)
    if fx is None:
        raise HTTPException(status_code=404, detail="Partita non trovata")
    kickoff = fx.kickoff if fx.kickoff is None or fx.kickoff.tzinfo else fx.kickoff.replace(tzinfo=timezone.utc)
    if kickoff is not None and kickoff <= _now():
        raise HTTPException(status_code=400, detail="La partita è già iniziata: non si può più aggiungere al diario")
    market_key = str(body.get("market_key") or "")
    market = MARKET_BY_KEY.get(market_key)
    if market is None:
        raise HTTPException(status_code=400, detail="Pronostico non valido")
    source = str(body.get("source") or "today")
    if source not in SOURCES:
        raise HTTPException(status_code=400, detail="Origine non valida")
    quota = _float(body.get("quota"), "Quota", minimum=1.01)
    if quota is None:
        raise HTTPException(status_code=400, detail="Inserisci la quota")
    duplicate = db.execute(
        select(AnalystEntry.id).where(
            AnalystEntry.user_id == user.id,
            AnalystEntry.today_fixture_id == fixture_id,
            AnalystEntry.market_key == market_key,
        )
    ).first()
    if duplicate:
        raise HTTPException(status_code=409, detail="Questo pronostico su questa partita è già nel tuo diario")
    cycle_id = body.get("cycle_id")
    if cycle_id:
        cycle = _own_cycle(db, principal, int(cycle_id))
        if cycle.status != CYCLE_OPEN:
            raise HTTPException(status_code=400, detail="La progressione è chiusa")
    entry = AnalystEntry(
        user_id=user.id,
        cycle_id=int(cycle_id) if cycle_id else None,
        **fixture_fields(fx),
        source=source,
        market_key=market_key,
        market_label=market["label"],
        quota=quota,
        stake=_float(body.get("stake"), "Puntata", minimum=0.01),
        stake_suggested=_float(body.get("stake_suggested"), "Puntata suggerita"),
        tool_json=tool_snapshot(fx, market_key, body.get("tool") if isinstance(body.get("tool"), dict) else None),
        checklist_json=clean_checklist(body.get("checklist")),
        analysis_note=(str(body.get("analysis_note") or "").strip() or None),
        notes=(str(body.get("notes") or "").strip() or None),
        minutes_spent=int(body["minutes_spent"]) if str(body.get("minutes_spent") or "").isdigit() else None,
        outcome=ENTRY_PENDING,
    )
    db.add(entry)
    db.commit()
    db.refresh(entry)
    return {"entry": entry_payload(entry)}


@router.patch("/entries/{entry_id}")
def update_entry(entry_id: int, body: dict[str, Any], principal: Principal = Depends(journal_principal), db: Session = Depends(get_db)) -> dict[str, Any]:
    entry = _own_entry(db, principal, entry_id)
    body = body or {}
    locked = is_locked(entry)
    if "notes" in body:
        entry.notes = str(body.get("notes") or "").strip() or None
    if "minutes_spent" in body:
        raw = body.get("minutes_spent")
        entry.minutes_spent = int(raw) if str(raw or "").isdigit() else None
    pre_match = {"market_key", "quota", "stake", "cycle_id", "checklist", "analysis_note", "stake_suggested"} & set(body)
    if pre_match and locked:
        raise HTTPException(status_code=400, detail="La partita è iniziata: si possono cambiare solo le note")
    if "market_key" in body:
        market = MARKET_BY_KEY.get(str(body["market_key"]))
        if market is None:
            raise HTTPException(status_code=400, detail="Pronostico non valido")
        entry.market_key, entry.market_label = market["key"], market["label"]
    if "quota" in body:
        quota = _float(body["quota"], "Quota", minimum=1.01)
        if quota is None:
            raise HTTPException(status_code=400, detail="Inserisci la quota")
        entry.quota = quota
    if "stake" in body:
        entry.stake = _float(body["stake"], "Puntata", minimum=0.01)
    if "stake_suggested" in body:
        entry.stake_suggested = _float(body["stake_suggested"], "Puntata suggerita")
    if "cycle_id" in body:
        cycle_id = body.get("cycle_id")
        if cycle_id:
            cycle = _own_cycle(db, principal, int(cycle_id))
            if cycle.status != CYCLE_OPEN:
                raise HTTPException(status_code=400, detail="La progressione è chiusa")
        entry.cycle_id = int(cycle_id) if cycle_id else None
    if "checklist" in body:
        entry.checklist_json = clean_checklist(body.get("checklist"))
    if "analysis_note" in body:
        entry.analysis_note = str(body.get("analysis_note") or "").strip() or None
    db.commit()
    return {"entry": entry_payload(entry)}


@router.delete("/entries/{entry_id}")
def delete_entry(entry_id: int, principal: Principal = Depends(journal_principal), db: Session = Depends(get_db)) -> dict[str, Any]:
    """Solo prima del calcio d'inizio (per correggere un errore): dopo, la giocata resta nello storico."""
    entry = _own_entry(db, principal, entry_id)
    if is_locked(entry):
        raise HTTPException(status_code=400, detail="La partita è iniziata: la giocata resta nel diario")
    db.delete(entry)
    db.commit()
    return {"deleted": entry_id}


# --- progressioni ---------------------------------------------------------------------------------


@router.get("/cycles")
def list_cycles(
    user_id: int | None = None,
    status: str | None = None,
    principal: Principal = Depends(journal_principal),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    scope = _scope(principal, user_id)
    settle_entries(db, user_id=scope)
    q = select(AnalystCycle).order_by(AnalystCycle.status.asc(), AnalystCycle.created_at.desc())
    if scope is not None:
        q = q.where(AnalystCycle.user_id == scope)
    if status:
        q = q.where(AnalystCycle.status == status)
    cycles = db.execute(q).scalars().all()
    names = _names(db, {c.user_id for c in cycles})
    return {"cycles": [{**cycle_payload(db, c), "user_name": names.get(c.user_id)} for c in cycles]}


@router.get("/cycles/{cycle_id}")
def get_cycle(cycle_id: int, quota: float | None = None, principal: Principal = Depends(journal_principal), db: Session = Depends(get_db)) -> dict[str, Any]:
    cycle = _visible_cycle(db, principal, cycle_id)
    settle_entries(db, user_id=cycle.user_id)
    return {"cycle": cycle_payload(db, cycle, quota=quota), "entries": [entry_payload(e) for e in cycle_entries(db, cycle)]}


@router.post("/cycles")
def create_cycle(body: dict[str, Any], principal: Principal = Depends(journal_principal), db: Session = Depends(get_db)) -> dict[str, Any]:
    user = _writer(principal, db)
    body = body or {}
    method = str(body.get("method") or "")
    if method not in METHODS:
        raise HTTPException(status_code=400, detail="Metodo non valido")
    try:
        settings = normalize_settings(method, body.get("settings") if isinstance(body.get("settings"), dict) else {})
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from None
    name = str(body.get("name") or "").strip()[:120] or f"{METHOD_LABELS[method]} {date.today().strftime('%d/%m/%Y')}"
    cycle = AnalystCycle(user_id=user.id, name=name, method=method, settings_json=settings, status=CYCLE_OPEN, notes=(str(body.get("notes") or "").strip() or None))
    db.add(cycle)
    db.commit()
    db.refresh(cycle)
    return {"cycle": cycle_payload(db, cycle)}


@router.patch("/cycles/{cycle_id}")
def update_cycle(cycle_id: int, body: dict[str, Any], principal: Principal = Depends(journal_principal), db: Session = Depends(get_db)) -> dict[str, Any]:
    cycle = _own_cycle(db, principal, cycle_id)
    body = body or {}
    if "name" in body:
        cycle.name = str(body.get("name") or "").strip()[:120] or cycle.name
    if "notes" in body:
        cycle.notes = str(body.get("notes") or "").strip() or None
    if "settings" in body:
        if any(e.outcome != ENTRY_PENDING for e in cycle_entries(db, cycle)):
            raise HTTPException(status_code=400, detail="Le impostazioni non si cambiano dopo la prima giocata chiusa")
        try:
            cycle.settings_json = normalize_settings(cycle.method, body.get("settings") or {})
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from None
    if "status" in body:
        status = str(body.get("status"))
        if status not in (CYCLE_OPEN, CYCLE_CLOSED):
            raise HTTPException(status_code=400, detail="Stato non valido")
        cycle.status = status
        cycle.closed_at = _now() if status == CYCLE_CLOSED else None
    db.commit()
    return {"cycle": cycle_payload(db, cycle)}


# --- report ed esportazione -----------------------------------------------------------------------


@router.get("/reports")
def reports(
    date_from: date | None = None,
    date_to: date | None = None,
    user_id: int | None = None,
    principal: Principal = Depends(journal_principal),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    settle_entries(db, user_id=_scope(principal, user_id))
    rows = db.execute(
        _entries_query(principal, user_id=user_id, date_from=date_from, date_to=date_to, cycle_id=None, outcome=None)
    ).scalars().all()
    return build_report(list(rows))


_OUTCOME_XLS = {"won": "VINTA", "lost": "PERSA", "void": "NULLA", "pending": ""}


@router.get("/export.xlsx")
def export_xlsx(
    date_from: date | None = None,
    date_to: date | None = None,
    user_id: int | None = None,
    principal: Principal = Depends(journal_principal),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    """Le giocate nello schema del foglio "GARE ATTUALE KW", piu' puntata, profitto e scheda."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    rows = list(
        db.execute(
            _entries_query(principal, user_id=user_id, date_from=date_from, date_to=date_to, cycle_id=None, outcome=None)
        ).scalars().all()
    )
    rows.reverse()  # dal piu' vecchio, come nel file
    names = _names(db, {e.user_id for e in rows})
    cycles = {c.id: c.name for c in db.execute(select(AnalystCycle)).scalars().all()} if rows else {}
    wb = Workbook()
    ws = wb.active
    ws.title = "GARE"
    headers = [
        "GARA", "DATA", "KW", "INCONTRO", "CAMPIONATO", "PRONOSTICO", "QUOTA", "RISULTATO PT", "RISULTATO FINALE",
        "ESITO", "VERSIONE SCELTA", "PUNTATA €", "PROFITTO €", "ORIGINE", "PROB. TOOL %", "MINUTI", "ANALISTA",
        "MODALITÀ DI INTERVENTO", "NOTE",
    ] + [c["label"] for c in CHECKLIST]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    for i, e in enumerate(rows, start=1):
        prob = (e.tool_json or {}).get("probability", (e.tool_json or {}).get("prob_cecchino"))
        try:
            prob_pct = round(float(prob) * (100 if float(prob) <= 1 else 1), 1) if prob is not None else None
        except (TypeError, ValueError):
            prob_pct = None
        checklist = e.checklist_json or {}
        ws.append(
            [
                i,
                e.match_date,
                e.match_date.isocalendar()[1],
                f"{e.home_team} - {e.away_team}",
                e.league_name,
                e.market_label,
                e.quota,
                e.result_ht,
                e.result_ft,
                _OUTCOME_XLS.get(e.outcome, e.outcome),
                cycles.get(e.cycle_id) if e.cycle_id else None,
                e.stake,
                e.profit,
                SOURCES.get(e.source, e.source),
                prob_pct,
                e.minutes_spent,
                names.get(e.user_id),
                e.analysis_note,
                e.notes,
            ]
            + [
                " · ".join(x for x in ((checklist.get(c["key"]) or {}).get("value"), (checklist.get(c["key"]) or {}).get("note")) if x) or None
                for c in CHECKLIST
            ]
        )
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
        row[1].number_format = "DD/MM/YYYY"
    for col, width in zip("ABCDEFGHIJKLMNOPQRS", (6, 11, 5, 34, 24, 13, 7, 11, 13, 9, 18, 10, 11, 15, 11, 8, 18, 40, 40)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"diario_analisi_{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
